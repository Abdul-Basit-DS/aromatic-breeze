"""
Orders admin – full enterprise order management system.
"""
from django.contrib import admin
from django.utils.html import format_html
from django.utils.translation import gettext_lazy as _
from django.http import HttpResponse
from django.urls import reverse
from django.utils import timezone
from import_export.admin import ExportMixin
from import_export import resources
import csv

from .models import Order, OrderItem, OrderStatusHistory, ShippingZone, ReturnRequest
from apps.accounts.models import ActivityLog


# ─── Resources ────────────────────────────────────────────────
class OrderResource(resources.ModelResource):
    class Meta:
        model = Order
        fields = (
            'order_number', 'customer_name', 'customer_email', 'customer_phone',
            'shipping_city', 'shipping_province', 'payment_method',
            'payment_status', 'status', 'subtotal', 'discount_amount',
            'shipping_charge', 'grand_total', 'coupon_code', 'created_at',
        )
        export_order = fields


# ─── Inlines ──────────────────────────────────────────────────
class OrderItemInline(admin.TabularInline):
    model = OrderItem
    extra = 0
    readonly_fields = ('product', 'variation', 'product_name', 'variation_name',
                       'sku', 'unit_price', 'quantity', 'line_total')
    can_delete = False

    def line_total(self, obj):
        return format_html('Rs. {}', obj.total_price)
    line_total.short_description = 'Total'


class OrderStatusHistoryInline(admin.TabularInline):
    model = OrderStatusHistory
    extra = 0
    readonly_fields = ('old_status', 'new_status', 'changed_by',
                       'note', 'customer_notified', 'created_at')
    can_delete = False
    ordering = ('-created_at',)


class ReturnRequestInline(admin.TabularInline):
    model = ReturnRequest
    extra = 0
    readonly_fields = ('reason', 'description', 'status', 'created_at')
    can_delete = False


# ─── Order Admin ──────────────────────────────────────────────
@admin.register(Order)
class OrderAdmin(ExportMixin, admin.ModelAdmin):
    resource_class = OrderResource

    list_display = (
        'order_number', 'customer_info', 'status_badge',
        'payment_badge', 'items_count', 'grand_total_display',
        'payment_method', 'created_at', 'actions_column'
    )
    list_filter = (
        'status', 'payment_status', 'payment_method', 'shipping_type',
        'created_at', 'shipping_city', 'shipping_province'
    )
    search_fields = (
        'order_number', 'customer_name', 'customer_email',
        'customer_phone', 'tracking_number', 'shipping_city'
    )
    date_hierarchy = 'created_at'
    ordering = ('-created_at',)
    list_per_page = 25
    readonly_fields = (
        'id', 'order_number', 'created_at', 'updated_at',
        'subtotal', 'grand_total'
    )

    fieldsets = (
        ('Order Info', {
            'fields': ('order_number', 'id', 'user', 'status', 'created_at', 'updated_at')
        }),
        ('Customer', {
            'fields': ('customer_name', 'customer_email', 'customer_phone')
        }),
        ('Shipping Address', {
            'fields': (
                'shipping_full_name', 'shipping_phone',
                'shipping_address_line1', 'shipping_address_line2',
                'shipping_city', 'shipping_province',
                'shipping_postal_code', 'shipping_country',
                'shipping_type'
            )
        }),
        ('Payment', {
            'fields': ('payment_method', 'payment_status',
                       'payment_reference', 'payment_date')
        }),
        ('Pricing', {
            'fields': (
                'subtotal', 'discount_amount', 'coupon_code',
                'coupon_discount', 'shipping_charge', 'cod_charge',
                'tax_amount', 'grand_total'
            )
        }),
        ('Tracking', {
            'fields': ('tracking_number', 'courier_name',
                       'shipped_at', 'delivered_at', 'estimated_delivery')
        }),
        ('Notes', {
            'fields': ('customer_notes', 'admin_notes')
        }),
    )

    inlines = [OrderItemInline, OrderStatusHistoryInline, ReturnRequestInline]

    actions = [
        'mark_confirmed', 'mark_packed', 'mark_shipped',
        'mark_delivered', 'mark_cancelled',
        'export_csv', 'export_excel',
    ]

    # ─── Display methods ──────────────────────────────────────
    def customer_info(self, obj):
        return format_html(
            '<strong>{}</strong><br>'
            '<small style="color:#888;">{}</small><br>'
            '<small>{}</small>',
            obj.customer_name, obj.customer_email, obj.customer_phone
        )
    customer_info.short_description = 'Customer'

    def status_badge(self, obj):
        color_map = {
            'pending': '#ffc107', 'confirmed': '#0dcaf0',
            'packed': '#0dcaf0', 'processing': '#0d6efd',
            'shipped': '#6f42c1', 'out_for_delivery': '#fd7e14',
            'delivered': '#198754', 'cancelled': '#dc3545',
            'returned': '#dc3545', 'refunded': '#6c757d',
        }
        color = color_map.get(obj.status, '#6c757d')
        return format_html(
            '<span style="background:{};color:white;padding:3px 10px;'
            'border-radius:20px;font-size:11px;font-weight:600;">{}</span>',
            color, obj.get_status_display()
        )
    status_badge.short_description = 'Status'

    def payment_badge(self, obj):
        colors = {
            'paid': '#198754', 'pending': '#ffc107',
            'failed': '#dc3545', 'refunded': '#6c757d', 'cod': '#0d6efd'
        }
        color = colors.get(obj.payment_status, '#6c757d')
        return format_html(
            '<span style="background:{};color:white;padding:2px 8px;'
            'border-radius:3px;font-size:11px;">{}</span>',
            color, obj.get_payment_status_display()
        )
    payment_badge.short_description = 'Payment'

    def items_count(self, obj):
        count = obj.items.count()
        return format_html('<strong>{}</strong> item{}', count, 's' if count != 1 else '')
    items_count.short_description = 'Items'

    def grand_total_display(self, obj):
        return format_html(
            '<strong style="color:#198754;">Rs. {}</strong>',
            obj.grand_total
    )
    grand_total_display.short_description = 'Total'

    def actions_column(self, obj):
        invoice_url = reverse('orders:invoice_pdf', kwargs={'order_number': obj.order_number})
        return format_html(
            '<a href="{}" target="_blank" style="'
            'background:#6c757d;color:white;padding:3px 8px;'
            'border-radius:3px;font-size:11px;text-decoration:none;">📄 Invoice</a>',
            invoice_url
        )
    actions_column.short_description = 'Actions'

    # ─── Save override – log status changes ───────────────────
    def save_model(self, request, obj, form, change):
        if change and 'status' in form.changed_data:
            old_status = form.initial.get('status', '')
            OrderStatusHistory.objects.create(
                order=obj,
                old_status=old_status,
                new_status=obj.status,
                changed_by=request.user,
                note=f'Status changed by {request.user.get_full_name()} via Admin'
            )
            ActivityLog.log(
                user=request.user,
                action=ActivityLog.ORDER_UPDATE,
                module='orders',
                description=f'Order {obj.order_number} status: {old_status} → {obj.status}',
                old_value={'status': old_status},
                new_value={'status': obj.status},
            )
        super().save_model(request, obj, form, change)

    # ─── Bulk status actions ───────────────────────────────────
    def _bulk_status(self, request, queryset, status, label):
        for order in queryset:
            old = order.status
            order.status = status
            order.save(update_fields=['status'])
            OrderStatusHistory.objects.create(
                order=order, old_status=old, new_status=status,
                changed_by=request.user,
                note=f'Bulk action: {label}'
            )
        self.message_user(request, f'{queryset.count()} orders marked as {label}.')

    def mark_confirmed(self, request, queryset):
        self._bulk_status(request, queryset, Order.CONFIRMED, 'Confirmed')
    mark_confirmed.short_description = 'Mark as Confirmed'

    def mark_packed(self, request, queryset):
        self._bulk_status(request, queryset, Order.PACKED, 'Packed')
    mark_packed.short_description = 'Mark as Packed'

    def mark_shipped(self, request, queryset):
        self._bulk_status(request, queryset, Order.SHIPPED, 'Shipped')
        queryset.filter(shipped_at__isnull=True).update(shipped_at=timezone.now())
    mark_shipped.short_description = 'Mark as Shipped'

    def mark_delivered(self, request, queryset):
        self._bulk_status(request, queryset, Order.DELIVERED, 'Delivered')
        queryset.filter(delivered_at__isnull=True).update(delivered_at=timezone.now())
    mark_delivered.short_description = 'Mark as Delivered'

    def mark_cancelled(self, request, queryset):
        cancellable = queryset.filter(status__in=[Order.PENDING, Order.CONFIRMED])
        self._bulk_status(request, cancellable, Order.CANCELLED, 'Cancelled')
    mark_cancelled.short_description = 'Cancel selected orders'

    def export_csv(self, request, queryset):
        response = HttpResponse(content_type='text/csv')
        response['Content-Disposition'] = 'attachment; filename="orders.csv"'
        writer = csv.writer(response)
        writer.writerow([
            'Order #', 'Date', 'Customer', 'Email', 'Phone',
            'City', 'Status', 'Payment', 'Method', 'Total'
        ])
        for o in queryset:
            writer.writerow([
                o.order_number, o.created_at.strftime('%Y-%m-%d %H:%M'),
                o.customer_name, o.customer_email, o.customer_phone,
                o.shipping_city, o.status, o.payment_status,
                o.payment_method, o.grand_total
            ])
        return response
    export_csv.short_description = 'Export to CSV'

    def export_excel(self, request, queryset):
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = 'Orders'
            headers = ['Order #', 'Date', 'Customer', 'Email', 'Phone',
                       'City', 'Province', 'Status', 'Payment Status',
                       'Payment Method', 'Subtotal', 'Discount',
                       'Shipping', 'Grand Total']
            header_fill = PatternFill('solid', fgColor='1a1a2e')
            for col, h in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=h)
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal='center')
            for row, o in enumerate(queryset, 2):
                ws.append([
                    o.order_number,
                    o.created_at.strftime('%Y-%m-%d %H:%M'),
                    o.customer_name, o.customer_email, o.customer_phone,
                    o.shipping_city, o.shipping_province,
                    o.get_status_display(), o.get_payment_status_display(),
                    o.get_payment_method_display(),
                    float(o.subtotal), float(o.discount_amount),
                    float(o.shipping_charge), float(o.grand_total)
                ])
            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                ws.column_dimensions[col[0].column_letter].width = max_len + 4
            response = HttpResponse(
                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            )
            response['Content-Disposition'] = 'attachment; filename="orders.xlsx"'
            wb.save(response)
            return response
        except ImportError:
            return self.export_csv(request, queryset)
    export_excel.short_description = 'Export to Excel'


# ─── Shipping Zone Admin ──────────────────────────────────────
@admin.register(ShippingZone)
class ShippingZoneAdmin(admin.ModelAdmin):
    list_display = ('name', 'standard_charge', 'express_charge',
                    'cod_charge', 'free_shipping_threshold', 'is_active')
    list_editable = ('standard_charge', 'express_charge', 'cod_charge', 'is_active')


# ─── Return Request Admin ─────────────────────────────────────
@admin.register(ReturnRequest)
class ReturnRequestAdmin(admin.ModelAdmin):
    list_display = ('id', 'order_link', 'reason', 'resolution_requested',
                    'status_badge', 'refund_amount', 'created_at')
    list_filter = ('status', 'reason', 'resolution_requested')
    search_fields = ('order__order_number', 'description')
    readonly_fields = ('order', 'order_item', 'reason', 'description',
                       'resolution_requested', 'created_at')
    list_per_page = 25

    fieldsets = (
        ('Request Info', {
            'fields': ('order', 'order_item', 'reason', 'description',
                       'resolution_requested', 'created_at')
        }),
        ('Admin Decision', {
            'fields': ('status', 'admin_notes', 'approved_by', 'refund_amount')
        }),
    )

    def order_link(self, obj):
        url = reverse('admin:orders_order_change', args=[obj.order.pk])
        return format_html('<a href="{}">#{}</a>', url, obj.order.order_number)
    order_link.short_description = 'Order'

    def status_badge(self, obj):
        colors = {
            'pending': '#ffc107', 'approved': '#198754',
            'rejected': '#dc3545', 'completed': '#6c757d'
        }
        color = colors.get(obj.status, '#6c757d')
        return format_html(
            '<span style="background:{};color:white;padding:2px 8px;'
            'border-radius:3px;font-size:11px;">{}</span>',
            color, obj.get_status_display()
        )
    status_badge.short_description = 'Status'
